# -*- coding: utf-8 -*-
"""포트폴리오 HTML 내용을 바탕으로 PowerPoint 생성 (원본 파일은 수정하지 않음)."""
import os
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor

# 경로
DESKTOP = Path(os.environ.get("USERPROFILE", "")) / "OneDrive" / "바탕 화면"
OUTPUT_DIR = DESKTOP / "포트폴리오_PPT"
PORTFOLIO = Path(__file__).resolve().parent
PROFILE_IMG = PORTFOLIO / "assets" / "profile.png"

PRIMARY = RGBColor(0x63, 0x66, 0xF1)
TEXT_DARK = RGBColor(0x1E, 0x29, 0x3B)
TEXT_LIGHT = RGBColor(0x64, 0x74, 0x8B)


def set_slide_title(slide, title: str, subtitle: str | None = None):
    if slide.shapes.title:
        slide.shapes.title.text = title
        tf = slide.shapes.title.text_frame
        tf.paragraphs[0].font.size = Pt(32)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = PRIMARY
    if subtitle and len(slide.placeholders) > 1:
        try:
            ph = slide.placeholders[1]
            ph.text = subtitle
            ph.text_frame.paragraphs[0].font.size = Pt(18)
            ph.text_frame.paragraphs[0].font.color.rgb = TEXT_LIGHT
        except Exception:
            pass


def add_bullet_slide(prs, title: str, bullets: list[str]):
    layout = prs.slide_layouts[1]  # Title and Content
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title
    slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = PRIMARY
    body = slide.placeholders[1].text_frame
    body.clear()
    for i, line in enumerate(bullets):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(18)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(8)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # --- 표지 ---
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    set_slide_title(slide, "강민지 포트폴리오", "열정적으로 배우고 성장하는 개발자 · 2026")

    # --- 자기소개 + 프로필 ---
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "About Me"
    slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = PRIMARY
    tf = slide.placeholders[1].text_frame
    tf.clear()
    lines = [
        "안녕하세요, 저는 강민지입니다.",
        "",
        "학부 과정에서 다양한 프로젝트를 경험하며 실무 역량을 키워왔습니다. 팀 프로젝트를 통해 협업의 중요성을 배웠고, 개인 프로젝트를 통해 자기주도적 학습 능력을 길렀습니다.",
        "",
        "• 이름: 강민지",
        "• 학력: 영남대학교 / 전자공학과",
        "• 이메일: kmj0313mj@naver.com",
    ]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(16)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(6)

    if PROFILE_IMG.is_file():
        left = Inches(9.2)
        top = Inches(2.0)
        slide.shapes.add_picture(str(PROFILE_IMG), left, top, width=Inches(2.8))

    # --- Skills ---
    add_bullet_slide(
        prs,
        "Skills",
        [
            "Frontend: HTML, CSS, JavaScript",
            "Backend: Python, Java",
            "Database: MySQL",
            "Tools: Git, GitHub",
        ],
    )

    # --- 프로젝트 목록 ---
    add_bullet_slide(
        prs,
        "Projects 개요",
        [
            "1. 실시간 운전자 졸음방지 모니터링 시스템 — Jetson Nano, Python, OpenCV, MediaPipe",
            "2. 교통표지판 인식 및 자율주행 제어 알고리즘 — MobileNetV2, PyTorch, Jetson Nano",
            "3. 반려동물 실시간 추적 및 실종 감지 시스템 — Raspberry Pi 5, TFLite, Gimbal, SMTP",
            "4. YOLOv11n 기반 실시간 침입 감지 CCTV — YOLOv11n, Raspberry Pi 5, PiCamera2, OpenCV",
        ],
    )

    projects = [
        (
            "프로젝트 1 — DMS",
            "실시간 운전자 졸음방지 모니터링 시스템",
            "Jetson Nano 기반 실시간 얼굴 인증 및 졸음 감지 통합 시스템",
            [
                "기간: 2025.09 - 2025.12",
                "역할: 임베디드 SW 개발",
                "팀: Capstone Design 팀",
                "태그: #임베디드_최적화 #컴퓨터비전 #H/W제어 #상태머신설계",
            ],
        ),
        (
            "프로젝트 2 — 자율주행 제어",
            "교통표지판 인식 및 자율주행 제어 알고리즘",
            "MobileNetV2 기반 경량 AI 모델 설계 및 실시간 하드웨어 제어 알고리즘 개발",
            [
                "기간: 2025.03 - 2025.06",
                "역할: 모델 설계 / 제어 알고리즘 개발",
                "팀: 엣지컴퓨팅설계 텀프로젝트",
                "태그: #MobileNetV2 #모델_경량화 #임베디드_제어 #PyTorch",
            ],
        ),
        (
            "프로젝트 3 — 반려동물 추적",
            "반려동물 실시간 추적 및 실종 감지 시스템",
            "라즈베리파이 5 기반 TFLite 객체 추적 및 Gimbal 제어, SMTP 자동 알림 시스템",
            [
                "기간: 2025.05 - 2025.09",
                "역할: 엣지 AI / 추적·알림 파이프라인 구현",
                "성과: 대한전자공학회 대구경북지부 AI·반도체·ICT 학술대회 논문·발표",
                "태그: #RaspberryPi5 #TFLite #Gimbal_Control #SMTP_Alert",
            ],
        ),
        (
            "프로젝트 4 — 침입 감지 CCTV",
            "YOLOv11n 기반 실시간 침입 감지 CCTV 시스템 설계 및 구현",
            "YOLOv11n 기반 실시간 객체 탐지 및 효율적 자원 관리를 위한 이벤트 구동형 보안 솔루션 (학사 논문)",
            [
                "기간: 2025.10 - 2025.12",
                "역할: 시스템 설계·구현, 실험·논문 작성 (1저)",
                "소속·지도: 영남대 전자공학부 AI/SW 트랙 · 서영석 교수",
                "태그: #YOLOv11n #RaspberryPi5 #PiCamera2 #ROI_침입 #이벤트_녹화",
            ],
        ),
    ]

    for sec_title, proj_title, subtitle, meta_lines in projects:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = sec_title
        slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = PRIMARY
        tf = slide.placeholders[1].text_frame
        tf.clear()
        first = True
        for line in [proj_title, "", subtitle, ""] + meta_lines:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.text = line
            p.font.size = Pt(15) if line else Pt(8)
            if line == proj_title:
                p.font.bold = True
                p.font.size = Pt(20)
            p.font.color.rgb = TEXT_DARK
            p.space_after = Pt(4)

    # --- Experience ---
    add_bullet_slide(
        prs,
        "Experience",
        [
            "[기간] — [경험 제목] @ [기관/회사명] — 경험에 대한 설명을 작성하세요.",
            "[기간] — [경험 제목] @ [기관/회사명] — 경험에 대한 설명을 작성하세요.",
            "[기간] — [경험 제목] @ [기관/회사명] — 경험에 대한 설명을 작성하세요.",
            "(웹 포트폴리오와 동일한 플레이스홀더입니다. 내용을 채우면 슬라이드도 함께 수정하세요.)",
        ],
    )

    # --- Contact ---
    add_bullet_slide(
        prs,
        "Contact",
        [
            "함께 일하고 싶으시다면 언제든 연락해주세요!",
            "",
            "Email: kmj0313mj@naver.com",
            "Phone: 010-4554-2609",
            "GitHub: https://github.com/kmj0313mj",
        ],
    )

    # --- 마무리 ---
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    set_slide_title(slide, "감사합니다", "질문이 있으시면 편하게 연락 주세요.")

    out_path = OUTPUT_DIR / "강민지_포트폴리오.pptx"
    prs.save(str(out_path))
    print(f"저장 완료: {out_path}")


if __name__ == "__main__":
    main()
